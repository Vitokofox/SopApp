from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from .schemas import SOPCreate, StepCreate


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", _clean(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", text).strip().rstrip(":")


def _cell_below_or_right(sheet, row: int, column: int) -> str:
    immediate_right = _clean(sheet.cell(row, column + 1).value)
    if immediate_right:
        return immediate_right
    below = _clean(sheet.cell(row + 1, column).value)
    if below:
        return below
    for offset in range(2, 8):
        right = _clean(sheet.cell(row, column + offset).value)
        if right:
            return right
    return ""


def _metadata(sheet) -> dict[str, str]:
    aliases = {
        "revision": "revision",
        "procedimiento estandar": "title",
        "titulo": "title",
        "nombre": "title",
        "elaborador": "elaborator",
        "area": "area",
        "identificador": "code",
        "codigo": "code",
        "jefe de area": "chief_area",
        "jefe masso": "chief_masso",
        "jefe sso": "chief_masso",
        "fecha publicacion": "publication_date",
        "ultima actualizacion": "last_update",
        "matriz de flexibilidad": "flexibility_matrix",
        "regla clave": "key_rule",
        "riesgos": "risks_text",
        "incidente": "risks_text",
        "medida de control": "controls_text",
        "epp": "epp_text",
        "herramientas a utilizar": "tools_text",
        "herramientas": "tools_text",
    }
    result: dict[str, str] = {}
    for row in sheet.iter_rows():
        for cell in row:
            key = aliases.get(_normalize(cell.value))
            if key and key not in result:
                value = _cell_below_or_right(sheet, cell.row, cell.column)
                if value:
                    result[key] = value
    return result


def _find_header(sheet, name: str) -> tuple[int, int] | None:
    target = _normalize(name)
    for row in sheet.iter_rows():
        for cell in row:
            if _normalize(cell.value) == target:
                return cell.row, cell.column
    return None


def _parse_steps(sheet) -> list[StepCreate]:
    step_header = _find_header(sheet, "Paso")
    description_header = _find_header(sheet, "Descripción")
    procedure_header = _find_header(sheet, "Procedimiento")
    symbol_header = _find_header(sheet, "Símbolo")
    if not step_header or not description_header or not procedure_header:
        return []
    header_row, step_column = step_header
    description_column = description_header[1]
    procedure_column = procedure_header[1]
    symbol_column = symbol_header[1] if symbol_header and symbol_header[0] == header_row else None
    steps: list[StepCreate] = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        raw_number = sheet.cell(row_number, step_column).value
        if not isinstance(raw_number, (int, float)) or int(raw_number) != raw_number:
            continue
        description = _clean(sheet.cell(row_number, description_column).value)
        procedure = _clean(sheet.cell(row_number, procedure_column).value)
        if not description and not procedure:
            continue
        symbol = _clean(sheet.cell(row_number, symbol_column).value) if symbol_column else ""
        risks = [] if _normalize(symbol) in {"", "no aplica", "n0 aplica"} else [symbol]
        steps.append(StepCreate(
            step_number=int(raw_number), description=description,
            procedure_text=procedure, risk_codes=risks,
        ))
    return steps


def parse_excel(content: bytes, filename: str) -> SOPCreate:
    try:
        book = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Excel inválido: {filename}") from exc
    candidates: list[tuple[int, object, dict[str, str]]] = []
    for sheet in book.worksheets:
        metadata = _metadata(sheet)
        score = int(bool(metadata.get("code"))) * 4 + int(bool(metadata.get("title"))) * 4 + len(_parse_steps(sheet))
        candidates.append((score, sheet, metadata))
    _, sheet, pairs = max(candidates, key=lambda item: item[0])
    steps = _parse_steps(sheet)
    code = pairs.get("code", "").strip()
    title = pairs.get("title", "").strip()
    if not code or not title:
        missing = " y ".join(name for name, value in (("código", code), ("nombre/título", title)) if not value)
        raise HTTPException(status_code=422, detail=f"El Excel no contiene {missing} reconocible")
    return SOPCreate(**{**pairs, "code": code, "title": title, "steps": steps})


def parse_pdf(content: bytes, filename: str) -> SOPCreate:
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"PDF inválido: {filename}") from exc
    if not text.strip():
        raise HTTPException(status_code=422, detail="El PDF no contiene texto extraíble")
    code_match = re.search(r"(?im)^\s*c[oó]digo\s*[:\-]\s*(.+)$", text)
    title_match = re.search(r"(?im)^\s*t[ií]tulo\s*[:\-]\s*(.+)$", text)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return SOPCreate(
        code=(code_match.group(1).strip() if code_match else Path(filename).stem),
        title=(title_match.group(1).strip() if title_match else lines[0]),
    )


def parse_document(content: bytes, filename: str) -> SOPCreate:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return parse_excel(content, filename)
    if suffix == ".pdf":
        return parse_pdf(content, filename)
    raise HTTPException(status_code=415, detail="Solo se permiten archivos .xlsx y .pdf")
